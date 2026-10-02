"""Build the final paper from paper_template.md.

Outputs (paper/):
  RAG_Healthcare_Final_Paper.md    resolved Markdown (numbered citations)
  RAG_Healthcare_Final_Paper.docx  IEEE-style two-column Word document

Layout follows the reference document's conventions (Letter, 24 pt centered title,
single-column abstract block, continuous break to two columns, small-caps section
headings, italic subsection headings, 8 pt captions, 8 pt hanging references).

Usage: python build_paper.py            (needs python-docx and lxml)
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
PAPER = HERE.parent
TEMPLATE = HERE / "paper_template.md"
FONT = "Times New Roman"

BODY_PT = float(sys.argv[1]) if len(sys.argv) > 1 else 10.0
LINE = float(sys.argv[2]) if len(sys.argv) > 2 else 1.0
COL_W = 3.487  # inches: (8.5 - 0.68 - 0.68 - 0.166) / 2

# --------------------------------------------------------------------------- references
REFS = {
    "lewis": 'P. Lewis <i>et al.</i>, "Retrieval-augmented generation for knowledge-intensive NLP tasks," in <i>Advances in Neural Information Processing Systems (NeurIPS)</i>, vol. 33, 2020, arXiv:2005.11401. doi: 10.48550/arXiv.2005.11401.',
    "xiong": 'G. Xiong, Q. Jin, Z. Lu, and A. Zhang, "Benchmarking retrieval-augmented generation for medicine," in <i>Findings of the Association for Computational Linguistics: ACL 2024</i>, 2024, pp. 6233–6251, doi: 10.18653/v1/2024.findings-acl.372.',
    "zhao": 'Y. Zhao, Y. Miao, R. Guo, Y. Luo, H. Wang, and Y. Wu, "Evaluation methods for inference-time retrieval-augmented and graph retrieval-augmented large language models in health care: Scoping review," <i>J. Med. Internet Res.</i>, vol. 28, p. e90046, 2026, doi: 10.2196/90046.',
    "joren": 'H. Joren, J. Zhang, C.-S. Ferng, D.-C. Juan, A. Taly, and C. Rashtchian, "Sufficient context: A new lens on retrieval augmented generation systems," in <i>Proc. Int. Conf. Learn. Represent. (ICLR)</i>, 2025, arXiv:2411.06037.',
    "meta": 'Meta AI, "Llama 3.2: Revolutionizing edge AI and vision with open, customizable models," blog post, Sep. 25, 2024. [Online]. Available: https://ai.meta.com/blog/llama-3-2-connect-2024-vision-edge-mobile-devices/',
    "robertson": 'S. Robertson and H. Zaragoza, "The probabilistic relevance framework: BM25 and beyond," <i>Found. Trends Inf. Retr.</i>, vol. 3, no. 4, pp. 333–389, 2009, doi: 10.1561/1500000019.',
    "reimers": 'N. Reimers and I. Gurevych, "Sentence-BERT: Sentence embeddings using Siamese BERT-networks," in <i>Proc. EMNLP-IJCNLP</i>, 2019, pp. 3982–3992, doi: 10.18653/v1/D19-1410.',
    "johnson": 'J. Johnson, M. Douze, and H. Jégou, "Billion-scale similarity search with GPUs," arXiv:1702.08734, 2017, doi: 10.48550/arXiv.1702.08734.',
    "cormack": 'G. V. Cormack, C. L. A. Clarke, and S. Büttcher, "Reciprocal rank fusion outperforms Condorcet and individual rank learning methods," in <i>Proc. 32nd Int. ACM SIGIR Conf. Res. Develop. Inf. Retr.</i>, 2009, pp. 758–759, doi: 10.1145/1571941.1572114.',
    "niu": 'C. Niu <i>et al.</i>, "RAGTruth: A hallucination corpus for developing trustworthy retrieval-augmented language models," in <i>Proc. 62nd Annu. Meeting Assoc. Comput. Linguistics (ACL)</i>, vol. 1, 2024, pp. 10862–10878, doi: 10.18653/v1/2024.acl-long.585.',
    "ma": 'Y. Ma, B. Chu, and N. Fuhr, "VERICITE: Evaluating sentence-level citation faithfulness in retrieval-augmented medical question answering," in <i>Proc. BioNLP 2026</i>, 2026, pp. 753–759, doi: 10.18653/v1/2026.bionlp-1.62.',
    "wilcoxon": 'F. Wilcoxon, "Individual comparisons by ranking methods," <i>Biometrics Bulletin</i>, vol. 1, no. 6, pp. 80–83, 1945, doi: 10.2307/3001968.',
    "holm": 'S. Holm, "A simple sequentially rejective multiple test procedure," <i>Scand. J. Statist.</i>, vol. 6, no. 2, pp. 65–70, 1979. [Online]. Available: https://www.jstor.org/stable/4615733',
}

# --------------------------------------------------------------------------- template parsing


def parse_template(text):
    blocks = []
    lines = text.splitlines()
    i = 0
    para = []

    def flush():
        nonlocal para
        if para:
            blocks.append(("p", " ".join(para)))
            para = []

    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            flush()
            i += 1
            continue
        if ln.startswith("# "):
            flush(); blocks.append(("title", ln[2:].strip())); i += 1; continue
        if ln.startswith("### "):
            flush(); blocks.append(("h2", ln[4:].strip())); i += 1; continue
        if ln.startswith("## "):
            flush(); blocks.append(("h1", ln[3:].strip())); i += 1; continue
        m = re.match(r"^(AUTHOR|AFFIL|EMAIL|ABSTRACT|KEYWORDS):\s*(.*)$", ln)
        if m:
            flush(); blocks.append((m.group(1).lower(), m.group(2).strip())); i += 1; continue
        if ln.startswith("$$") and ln.endswith("$$") and len(ln) > 4:
            flush(); blocks.append(("eq", ln[2:-2].strip())); i += 1; continue
        m = re.match(r"^!\[(.*)\]\((.*?)\)(?:\{w=([\d.]+)\})?$", ln)
        if m:
            flush(); blocks.append(("fig", m.group(1), m.group(2), float(m.group(3) or 3.0))); i += 1; continue
        if ln.startswith("Table:"):
            flush()
            cap = ln[len("Table:"):].strip()
            opts = {}
            m2 = re.search(r"\{([^}]*)\}\s*$", cap)
            if m2:
                for part in m2.group(1).split():
                    k, _, v = part.partition("=")
                    opts[k] = v
                cap = cap[:m2.start()].strip()
            rows = []
            i += 1
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", cap, rows, opts))
            continue
        para.append(ln.strip())
        i += 1
    flush()
    return blocks


def resolve_citations(blocks):
    order = []

    def repl(m):
        keys = [k.strip().lstrip("@") for k in m.group(1).split(",")]
        out = []
        for k in keys:
            if k not in REFS:
                raise KeyError(f"unknown citation key {k}")
            if k not in order:
                order.append(k)
            out.append(f"[{order.index(k) + 1}]")
        return ", ".join(out)

    pat = re.compile(r"\[(@[a-z]+(?:\s*,\s*@[a-z]+)*)\]")
    new = []
    for b in blocks:
        b = list(b)
        for idx, v in enumerate(b):
            if isinstance(v, str) and idx > 0 and b[0] in ("p", "abstract", "keywords", "fig"):
                b[idx] = pat.sub(repl, v) if idx == 1 or (b[0] == "fig" and idx == 1) else v
        if b[0] == "table":
            b[1] = pat.sub(repl, b[1])
        new.append(tuple(b))
    unused = [k for k in REFS if k not in order]
    if unused:
        raise SystemExit(f"uncited references: {unused}")
    return new, order


# --------------------------------------------------------------------------- LaTeX -> OMML
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
SYMS = {"cdot": "⋅", "cap": "∩", "in": "∈", "le": "≤", "subseteq": "⊆",
        "sqsubseteq": "⊑", "nu": "ν", "{": "{", "}": "}"}


def tokenize(s):
    toks = []
    i = 0
    while i < len(s):
        c = s[i]
        if c == "\\":
            j = i + 1
            if j < len(s) and s[j].isalpha():
                while j < len(s) and s[j].isalpha():
                    j += 1
                toks.append(s[i:j])
            else:
                toks.append(s[i:j + 1])
                j += 1
            i = j
        elif c.isspace():
            i += 1
        else:
            toks.append(c)
            i += 1
    return toks


def esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def mrun(text, sty=None, scr=None):
    pr = ""
    if sty or scr:
        pr = "<m:rPr>" + (f'<m:scr m:val="{scr}"/>' if scr else "") + (f'<m:sty m:val="{sty}"/>' if sty else "") + "</m:rPr>"
    return (f'<m:r>{pr}<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/><w:sz w:val="20"/></w:rPr>'
            f"<m:t xml:space=\"preserve\">{esc(text)}</m:t></m:r>")


class MathParser:
    def __init__(self, toks):
        self.t = toks
        self.p = 0

    def peek(self):
        return self.t[self.p] if self.p < len(self.t) else None

    def group(self):
        assert self.peek() == "{", self.peek()
        self.p += 1
        out = self.seq(stop="}")
        assert self.peek() == "}"
        self.p += 1
        return out

    def raw_group(self):
        assert self.peek() == "{"
        self.p += 1
        buf = []
        while self.peek() != "}":
            buf.append(self.t[self.p]); self.p += 1
        self.p += 1
        return "".join(buf)

    def seq(self, stop=None):
        out = []
        while self.peek() is not None and self.peek() != stop and self.peek() != "}":
            out.append(self.atom_with_scripts())
        return "".join(out)

    def atom_with_scripts(self):
        tok = self.peek()
        if tok == "\\sum":
            self.p += 1
            sub = sup = ""
            while self.peek() in ("_", "^"):
                kind = self.t[self.p]; self.p += 1
                g = self.group() if self.peek() == "{" else self.atom_with_scripts()
                if kind == "_":
                    sub = g
                else:
                    sup = g
            body = self.seq(stop=None)
            return ('<m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/>'
                    + ("" if sub else '<m:subHide m:val="1"/>') + ("" if sup else '<m:supHide m:val="1"/>')
                    + f"</m:naryPr><m:sub>{sub}</m:sub><m:sup>{sup}</m:sup><m:e>{body}</m:e></m:nary>")
        base = self.atom()
        sub = sup = None
        while self.peek() in ("_", "^"):
            kind = self.t[self.p]; self.p += 1
            g = self.group() if self.peek() == "{" else self.atom()
            if kind == "_":
                sub = g
            else:
                sup = g
        if sub is not None and sup is not None:
            return f"<m:sSubSup><m:e>{base}</m:e><m:sub>{sub}</m:sub><m:sup>{sup}</m:sup></m:sSubSup>"
        if sub is not None:
            return f"<m:sSub><m:e>{base}</m:e><m:sub>{sub}</m:sub></m:sSub>"
        if sup is not None:
            return f"<m:sSup><m:e>{base}</m:e><m:sup>{sup}</m:sup></m:sSup>"
        return base

    def atom(self):
        tok = self.t[self.p]
        self.p += 1
        if tok == "\\frac":
            num = self.group()
            den = self.group()
            return f"<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>"
        if tok in ("\\mathrm", "\\text"):
            return mrun(self.raw_group(), sty="p")
        if tok == "\\mathbf":
            return mrun(self.raw_group(), sty="b")
        if tok == "\\mathbb":
            return mrun(self.raw_group(), sty="p", scr="double-struck")
        if tok in ("\\log", "\\min"):
            return mrun(tok[1:], sty="p")
        if tok == "\\ ":
            return mrun(" ", sty="p")
        if tok.startswith("\\"):
            name = tok[1:]
            if name in SYMS:
                return mrun(SYMS[name], sty="p")
            raise ValueError(f"unsupported command {tok}")
        if tok == "{":
            self.p -= 1
            return self.group()
        if tok.isalpha():
            return mrun(tok)
        return mrun(tok, sty="p")


def latex_to_omml(latex):
    body = MathParser(tokenize(latex)).seq()
    return (f'<m:oMathPara xmlns:m="{M_NS}" xmlns:w="{W_NS}"><m:oMathParaPr><m:jc m:val="center"/></m:oMathParaPr>'
            f"<m:oMath>{body}</m:oMath></m:oMathPara>")


# --------------------------------------------------------------------------- docx helpers
def set_run_font(run, size=None, bold=None, italic=None, small_caps=None, name=FONT, color=None):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = parse_xml(f'<w:rFonts xmlns:w="{W_NS}"/>')
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if small_caps is not None:
        run.font.small_caps = small_caps
    run.font.color.rgb = RGBColor(0, 0, 0) if color is None else color


def add_inline(par, text, size, base_bold=False, base_italic=False):
    """Inline markup: **bold**, *italic*, <i>..</i>, _{subscript}, \\* literal asterisk."""
    text = text.replace("\\*", "\u0001")
    pat = re.compile(r"(\*\*.+?\*\*|\*.+?\*|<i>.+?</i>|_\{[^}]*\})")
    for tok in pat.split(text):
        if not tok:
            continue
        bold, italic, sub = base_bold, base_italic, False
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            tok, bold = tok[2:-2], True
        elif tok.startswith("<i>"):
            tok, italic = tok[3:-4], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, italic = tok[1:-1], True
        elif tok.startswith("_{"):
            tok, sub = tok[2:-1], True
        r = par.add_run(tok.replace("\u0001", "*"))
        set_run_font(r, size=size, bold=bold, italic=italic)
        if sub:
            r.font.subscript = True


def fmt(par, align=None, before=0, after=0, first=None, left=None, line=None, keep_next=False, keep_lines=False):
    pf = par.paragraph_format
    if align is not None:
        par.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if first is not None:
        pf.first_line_indent = Pt(first)
    if left is not None:
        pf.left_indent = Pt(left)
    if line is not None:
        pf.line_spacing = line
    pf.keep_with_next = keep_next
    pf.keep_together = keep_lines
    pf.widow_control = True


def set_cell_borders(table, sz=4, color="000000"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders xmlns:w="{W_NS}">'
        + "".join(f'<w:{e} w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                  for e in ("top", "left", "bottom", "right", "insideH", "insideV"))
        + "</w:tblBorders>")
    tblPr.append(borders)


def no_borders(table):
    tblPr = table._tbl.tblPr
    tblPr.append(parse_xml(
        f'<w:tblBorders xmlns:w="{W_NS}">'
        + "".join(f'<w:{e} w:val="nil"/>' for e in ("top", "left", "bottom", "right", "insideH", "insideV"))
        + "</w:tblBorders>"))


def fix_table_layout(table, widths_in, margins_twips=(15, 28)):
    tblPr = table._tbl.tblPr
    total = int(sum(widths_in) * 1440)
    for tag in ("w:tblW", "w:tblLayout", "w:tblCellMar"):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    tblPr.append(parse_xml(f'<w:tblW xmlns:w="{W_NS}" w:w="{total}" w:type="dxa"/>'))
    tblPr.append(parse_xml(f'<w:tblLayout xmlns:w="{W_NS}" w:type="fixed"/>'))
    tv, lh = margins_twips
    tblPr.append(parse_xml(
        f'<w:tblCellMar xmlns:w="{W_NS}"><w:top w:w="{tv}" w:type="dxa"/><w:left w:w="{lh}" w:type="dxa"/>'
        f'<w:bottom w:w="{tv}" w:type="dxa"/><w:right w:w="{lh}" w:type="dxa"/></w:tblCellMar>'))
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_in):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in):
            cell.width = Inches(w)


def row_props(row, cant_split=True, header=False):
    trPr = row._tr.get_or_add_trPr()
    if cant_split:
        trPr.append(parse_xml(f'<w:cantSplit xmlns:w="{W_NS}"/>'))
    if header:
        trPr.append(parse_xml(f'<w:tblHeader xmlns:w="{W_NS}"/>'))


# --------------------------------------------------------------------------- build
def build():
    blocks = parse_template(TEMPLATE.read_text(encoding="utf-8"))
    blocks, cite_order = resolve_citations(blocks)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY_PT)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for nm, size in (("Heading 1", 10), ("Heading 2", 10), ("Heading 3", 8)):
        hs = doc.styles[nm]
        hs.font.name = FONT
        hs.font.size = Pt(size)
        hs.font.bold = False
        hs.font.color.rgb = RGBColor(0, 0, 0)
        rpr = hs.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = parse_xml(f'<w:rFonts xmlns:w="{W_NS}"/>')
            rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]

    core = doc.core_properties
    core.author = "Hatim Maula"
    core.title = "Reliability-Oriented Evaluation of a Local Healthcare-Document RAG System"
    core.subject = "Retrieval-augmented generation; healthcare documents; failure analysis"
    core.comments = ""

    sec0 = doc.sections[0]
    sec0.page_width, sec0.page_height = Inches(8.5), Inches(11)
    sec0.left_margin, sec0.right_margin = Inches(0.919), Inches(0.68)
    sec0.top_margin, sec0.bottom_margin = Inches(0.764), Inches(1.019)

    counters = {"fig": 0, "tab": 0, "eq": 0}
    roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
    md = []
    two_col_started = False

    def start_two_columns():
        nonlocal two_col_started
        sec = doc.add_section(WD_SECTION.CONTINUOUS)
        sec.left_margin, sec.right_margin = Inches(0.68), Inches(0.68)
        sec.top_margin, sec.bottom_margin = Inches(0.702), Inches(1.019)
        sectPr = sec._sectPr
        for c in sectPr.findall(qn("w:cols")):
            sectPr.remove(c)
        sectPr.find(qn("w:pgMar")).addnext(parse_xml(f'<w:cols xmlns:w="{W_NS}" w:num="2" w:space="239"/>'))
        two_col_started = True

    for blk in blocks:
        kind = blk[0]
        if kind == "title":
            p = doc.add_paragraph()
            fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 6, line=0.95)
            r = p.add_run(blk[1]); set_run_font(r, size=24)
            md.append(f"# {blk[1]}\n")
        elif kind == "author":
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 0)
            set_run_font(p.add_run(blk[1]), size=12)
            md.append(f"**{blk[1]}**  ")
        elif kind == "affil":
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 0)
            set_run_font(p.add_run(blk[1]), size=10, italic=True)
            md.append(f"*{blk[1]}*  ")
        elif kind == "email":
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 8)
            set_run_font(p.add_run(blk[1]), size=9)
            md.append(f"{blk[1]}\n")
        elif kind in ("abstract", "keywords"):
            p = doc.add_paragraph()
            fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 3 if kind == "abstract" else 6, first=0, line=0.95)
            label = "Abstract—" if kind == "abstract" else "Keywords—"
            set_run_font(p.add_run(label), size=9, bold=True, italic=True)
            add_inline(p, blk[1], 9, base_bold=True)
            md.append(f"***{label}***{blk[1]}\n")
            if kind == "keywords":
                start_two_columns()
        elif kind == "h1":
            idx = None
            text = blk[1]
            p = doc.add_paragraph(style="Heading 1")
            fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 4, 1, first=0, left=0, keep_next=True)
            m = re.match(r"^([IVX]+)\.\s+(.*)$", text)
            label, title = m.group(1) + ".", m.group(2)
            set_run_font(p.add_run(label + " "), size=10, small_caps=False)
            set_run_font(p.add_run(title), size=10, small_caps=True)
            md.append(f"\n## {text}\n")
        elif kind == "h2":
            p = doc.add_paragraph(style="Heading 2")
            fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 2, 1, first=0, left=0, keep_next=True)
            set_run_font(p.add_run(blk[1]), size=10, italic=True)
            md.append(f"\n### {blk[1]}\n")
        elif kind == "p":
            p = doc.add_paragraph()
            fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 1, first=9.45, line=LINE)
            add_inline(p, blk[1], BODY_PT)
            md.append(blk[1].replace("\\*", "\\*") + "\n")
        elif kind == "eq":
            counters["eq"] += 1
            n = counters["eq"]
            t = doc.add_table(rows=1, cols=2)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            fix_table_layout(t, [COL_W - 0.38, 0.38], margins_twips=(0, 0))
            no_borders(t)
            row_props(t.rows[0])
            c0, c1 = t.rows[0].cells
            p0 = c0.paragraphs[0]
            fmt(p0, WD_ALIGN_PARAGRAPH.CENTER, 1, 1, first=0, line=1.0, keep_next=False)
            p0._p.append(parse_xml(latex_to_omml(blk[1])))
            p1 = c1.paragraphs[0]
            fmt(p1, WD_ALIGN_PARAGRAPH.RIGHT, 1, 1, first=0, line=1.0)
            set_run_font(p1.add_run(f"({n})"), size=BODY_PT)
            for c in (c0, c1):
                tcPr = c._tc.get_or_add_tcPr()
                tcPr.append(parse_xml(f'<w:vAlign xmlns:w="{W_NS}" w:val="center"/>'))
            md.append(f"$$ {blk[1]} \\tag{{{n}}} $$\n")
        elif kind == "fig":
            counters["fig"] += 1
            _, cap, path, width = blk
            p = doc.add_paragraph()
            fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 3, 1, first=0, keep_next=True)
            p.add_run().add_picture(str(PAPER / path), width=Inches(width))
            cp = doc.add_paragraph()
            fmt(cp, WD_ALIGN_PARAGRAPH.CENTER, 0, 6, first=0, line=1.0, keep_lines=True)
            add_inline(cp, cap, 8)
            md.append(f"![{cap}]({path})\n")
        elif kind == "table":
            counters["tab"] += 1
            _, cap, rows, opts = blk
            n = counters["tab"]
            widths = [float(x) for x in opts["w"].split(",")] if "w" in opts else None
            fs = float(opts.get("fs", 6.8))
            keep = opts.get("keep", "1") != "0"
            romans = ["I", "II", "III", "IV", "V", "VI"]
            p1 = doc.add_paragraph()
            fmt(p1, WD_ALIGN_PARAGRAPH.CENTER, 4, 0, first=0, keep_next=True)
            set_run_font(p1.add_run(f"TABLE {romans[n - 1]}"), size=8, small_caps=True)
            # first sentence of the caption is the title (small caps); the rest is a note
            title, _, note = cap.partition(". ")
            p2 = doc.add_paragraph()
            fmt(p2, WD_ALIGN_PARAGRAPH.CENTER, 0, 3, first=0, keep_next=True, line=1.0)
            set_run_font(p2.add_run(title.rstrip(".")), size=8, small_caps=True)
            ncols = len(rows[0])
            t = doc.add_table(rows=len(rows), cols=ncols)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            fix_table_layout(t, widths or [COL_W / ncols] * ncols)
            set_cell_borders(t)
            for ri, row in enumerate(rows):
                row_props(t.rows[ri], header=(ri == 0))
                for ci, val in enumerate(row):
                    cell = t.rows[ri].cells[ci]
                    cp = cell.paragraphs[0]
                    left_align = (ncols == 4 and widths and widths[1] > 1.0 and ci in (1, 3)) or ci == 0 and ncols != 5 and widths and widths[0] > 0.8
                    fmt(cp, WD_ALIGN_PARAGRAPH.LEFT if left_align else WD_ALIGN_PARAGRAPH.CENTER, 0, 0, first=0, line=1.0,
                        keep_next=(keep and ri < len(rows) - 1) or (ri == 0))
                    r = cp.add_run(val.replace("\\*", "*"))
                    set_run_font(r, size=(float(opts.get("hfs", fs)) if ri == 0 else fs), bold=(ri == 0))
            if note:
                np_ = doc.add_paragraph()
                fmt(np_, WD_ALIGN_PARAGRAPH.JUSTIFY, 2, 6, first=0, line=1.0)
                add_inline(np_, note, 7)
            else:
                sp = doc.add_paragraph(); fmt(sp, None, 0, 4, first=0)
                sp.paragraph_format.line_spacing = Pt(2)
            md.append(f"\n**TABLE {romans[n - 1]}.** {cap}\n")
            md.append("| " + " | ".join(rows[0]) + " |")
            md.append("|" + "|".join("---" for _ in rows[0]) + "|")
            for row in rows[1:]:
                md.append("| " + " | ".join(row) + " |")
            md.append("")

    # references
    hp = doc.add_paragraph(style="Heading 2")
    fmt(hp, WD_ALIGN_PARAGRAPH.CENTER, 6, 4, first=0, left=0, keep_next=True)
    set_run_font(hp.add_run("References"), size=10, small_caps=True, italic=False)
    md.append("\n## REFERENCES\n")
    for i, key in enumerate(cite_order, start=1):
        rp = doc.add_paragraph()
        fmt(rp, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 2, first=-22, left=22, line=1.0)
        rp.paragraph_format.tab_stops.add_tab_stop(Pt(22))
        set_run_font(rp.add_run(f"[{i}]\t"), size=8)
        add_inline(rp, REFS[key], 8)
        plain = re.sub(r"</?i>", "*", REFS[key])
        md.append(f"[{i}] {plain}\n")

    order = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
             "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
             "tblCaption", "tblDescription"]
    for t in doc.tables:
        pr = t._tbl.tblPr
        kids = list(pr)
        for k in kids:
            pr.remove(k)
        kids.sort(key=lambda e: order.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in order else 99)
        for k in kids:
            pr.append(k)
    for z in doc.settings.element.findall(qn("w:zoom")):
        if z.get(qn("w:percent")) is None:
            z.set(qn("w:percent"), "100")
    out_docx = PAPER / "RAG_Healthcare_Final_Paper.docx"
    doc.save(out_docx)
    (PAPER / "RAG_Healthcare_Final_Paper.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    print("built", out_docx.name, "citations:", len(cite_order), counters)


if __name__ == "__main__":
    build()
